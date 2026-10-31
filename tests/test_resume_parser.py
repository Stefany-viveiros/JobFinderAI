from docx import Document
from io import BytesIO

from jobfinder.resume_parser import parse_resume, extract_resume_text


def make_docx():
    doc = Document()
    doc.add_paragraph("Marina da Silva")
    doc.add_paragraph("marina@example.com | (11) 91234-5678")
    doc.add_paragraph("FORMAÇÃO ACADÊMICA")
    doc.add_paragraph("Análise e Desenvolvimento de Sistemas | Cruzeiro do Sul")
    doc.add_paragraph("Engenharia da Computação | FAM")
    doc.add_paragraph("HABILIDADES TÉCNICAS")
    doc.add_paragraph("Python, JavaScript, React e SQL")
    doc.add_paragraph("CURSOS COMPLEMENTARES")
    doc.add_paragraph("Python Básico | SENAI")
    doc.add_paragraph("Git e GitHub | Google")
    doc.add_paragraph("EXPERIÊNCIA PROFISSIONAL")
    doc.add_paragraph("Estágio de suporte técnico - Empresa Exemplo")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def test_parse_docx_extracts_profile_suggestions():
    parsed = parse_resume(make_docx(), "curriculo.docx")
    assert parsed["name"] == "Marina da Silva"
    assert "Python" in parsed["technologies"]
    assert "Análise e Desenvolvimento" in parsed["course"]
    assert "Estágio de suporte" in parsed["experience"]


def test_basic_parser_keeps_multiple_education_records_separate():
    parsed = parse_resume(make_docx(), "curriculo.docx")
    assert len(parsed["educations"]) == 2
    assert parsed["educations"][0]["course"] == "Análise e Desenvolvimento de Sistemas"
    assert parsed["educations"][0]["institution"] == "Cruzeiro do Sul"
    assert parsed["educations"][1]["course"] == "Engenharia da Computação"
    assert parsed["educations"][1]["institution"] == "FAM"


def test_basic_parser_keeps_courses_as_separate_records():
    parsed = parse_resume(make_docx(), "curriculo.docx")
    assert len(parsed["courses"]) == 2
    assert parsed["courses"][0]["name"] == "Python Básico"
    assert parsed["courses"][0]["provider"] == "SENAI"


def test_rejects_unsupported_format():
    try:
        parse_resume(b"texto qualquer", "curriculo.txt")
    except ValueError as exc:
        assert "PDF ou DOCX" in str(exc)
    else:
        raise AssertionError("Formato não suportado deveria ser rejeitado")



def test_extract_resume_text_keeps_full_docx_content():
    text = extract_resume_text(make_docx(), "curriculo.docx")
    assert "Marina da Silva" in text
    assert "FORMAÇÃO ACADÊMICA" in text
    assert "Estágio de suporte técnico" in text
